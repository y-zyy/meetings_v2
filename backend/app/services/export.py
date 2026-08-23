"""Generate DOCX, PDF, and TXT exports from a meeting record."""

import html as _html
import io
import os
import re
from datetime import date

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import RGBColor, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

# ── Font setup ───────────────────────────────────────────────────────────────
# Noto Sans KR must be embedded as real TTF files (ReportLab has no built-in
# CID font for it — that's only true of things like HYGothic/HYSMyeongJo).
# Put the .ttf files next to this script in a "fonts/" folder, e.g.:
#   fonts/NotoSansKR-Regular.ttf
#   fonts/NotoSansKR-Bold.ttf
# You can get these from Google Fonts: https://fonts.google.com/noto/specimen/Noto+Sans+KR

_FONT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fonts")

_KO_FONT = "NotoSansKR"
_KO_FONT_BOLD = "NotoSansKR-Bold"

pdfmetrics.registerFont(TTFont(_KO_FONT, os.path.join(_FONT_DIR, "NotoSansKR-Regular.ttf")))
pdfmetrics.registerFont(TTFont(_KO_FONT_BOLD, os.path.join(_FONT_DIR, "NotoSansKR-Bold.ttf")))
# Noto Sans KR has no dedicated italic cut, so map italic/bold-italic to the
# closest available weights rather than leaving them unregistered (which
# would make <b>/<i> tags silently fall back to Helvetica).
pdfmetrics.registerFontFamily(
    _KO_FONT,
    normal=_KO_FONT,
    bold=_KO_FONT_BOLD,
    italic=_KO_FONT,
    boldItalic=_KO_FONT_BOLD,
)

# DOCX-side font name (this is just a name Word resolves at render time —
# no embedding happens here unless you also embed the font in the .docx,
# so the font should be installed on whatever machine opens the file).
_DOCX_FONT = "Noto Sans KR"

# ── helpers ───────────────────────────────────────────────────────────────────

def _fmt_date(d: date | None) -> str:
    return d.strftime("%Y-%m-%d") if d else "미지정"


def _esc(text: str | None) -> str:
    """HTML-escape text for use in ReportLab Paragraph (XML parser)."""
    return _html.escape(str(text or ""))


def _esc_lines(text: str | None) -> str:
    """HTML-escape and convert newlines to <br/> for multi-line Paragraph."""
    return _esc(text).replace("\n", "<br/>")


def _fmt_duration(secs: int | None) -> str:
    if not secs:
        return "-"
    h, r = divmod(secs, 3600)
    m, s = divmod(r, 60)
    return f"{h}:{m:02d}:{s:02d}"


# ── Markdown parser ───────────────────────────────────────────────────────────

def _inline_to_plain(text: str) -> str:
    """Strip markdown inline markup to plain text."""
    text = re.sub(r'\*\*\*(.+?)\*\*\*', r'\1', text)
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'\*(.+?)\*', r'\1', text)
    text = re.sub(r'~~(.+?)~~', r'\1', text)
    text = re.sub(r'`([^`]+)`', r'\1', text)
    text = re.sub(r'!\[([^\]]*)\]\([^)]+\)', r'\1', text)
    text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)
    return text


def _inline_to_pdf(text: str) -> str:
    """Convert markdown inline markup to ReportLab XML tags."""
    text = _esc(text)
    text = re.sub(r'\*\*\*(.+?)\*\*\*', r'<b><i>\1</i></b>', text)
    text = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', text)
    text = re.sub(r'\*(.+?)\*', r'<i>\1</i>', text)
    text = re.sub(r'~~(.+?)~~', r'<strike>\1</strike>', text)
    text = re.sub(r'`([^`]+)`', r'<font face="Courier">\1</font>', text)
    text = re.sub(r'!\[([^\]]*)\]\([^)]+\)', r'\1', text)
    text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)
    return text


def _parse_markdown(src: str) -> list[dict]:
    """Parse markdown into a list of block dicts."""
    lines = src.replace('\r\n', '\n').replace('\r', '\n').split('\n')
    blocks: list[dict] = []
    i = 0
    ul_items: list[str] = []
    ol_items: list[str] = []

    def flush_lists():
        nonlocal ul_items, ol_items
        if ul_items:
            blocks.append({"type": "ul", "items": ul_items[:]})
            ul_items = []
        if ol_items:
            blocks.append({"type": "ol", "items": ol_items[:]})
            ol_items = []

    while i < len(lines):
        line = lines[i]

        # Fenced code block
        if line.startswith('```'):
            flush_lists()
            buf = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                buf.append(lines[i])
                i += 1
            blocks.append({"type": "code", "text": '\n'.join(buf)})
            i += 1
            continue

        # HR
        if re.match(r'^(-{3,}|\*{3,}|_{3,})\s*$', line):
            flush_lists()
            blocks.append({"type": "hr"})
            i += 1
            continue

        # Heading
        hm = re.match(r'^(#{1,6})\s+(.+)', line)
        if hm:
            flush_lists()
            blocks.append({"type": "heading", "level": len(hm.group(1)), "text": hm.group(2)})
            i += 1
            continue

        # Table (pipe-delimited)
        if line.startswith('|'):
            flush_lists()
            table_lines = []
            while i < len(lines) and lines[i].startswith('|'):
                table_lines.append(lines[i])
                i += 1
            if len(table_lines) >= 2 and re.match(r'^\|[\s\-:|]+\|', table_lines[1]):
                def parse_row(row):
                    return [c.strip() for c in re.sub(r'^\||\|$', '', row).split('|')]
                headers = parse_row(table_lines[0])
                rows = [parse_row(r) for r in table_lines[2:]]
                blocks.append({"type": "table", "headers": headers, "rows": rows})
            else:
                for tl in table_lines:
                    blocks.append({"type": "paragraph", "text": tl})
            continue

        # Unordered list
        if re.match(r'^[\*\-\+]\s+', line):
            if ol_items:
                blocks.append({"type": "ol", "items": ol_items[:]})
                ol_items = []
            ul_items.append(re.sub(r'^[\*\-\+]\s+', '', line))
            i += 1
            continue

        # Ordered list
        if re.match(r'^\d+\.\s+', line):
            if ul_items:
                blocks.append({"type": "ul", "items": ul_items[:]})
                ul_items = []
            ol_items.append(re.sub(r'^\d+\.\s+', '', line))
            i += 1
            continue

        flush_lists()

        # Blank line
        if not line.strip():
            i += 1
            continue

        # Paragraph
        blocks.append({"type": "paragraph", "text": line})
        i += 1

    flush_lists()
    return blocks


# ── DOCX ─────────────────────────────────────────────────────────────────────

def _docx_set_style_font(doc, font_name: str = _DOCX_FONT):
    """Set the Latin + East-Asian font on the styles used throughout the doc.

    Setting run.font.name alone is not enough for Korean text in Word —
    Word keeps a separate 'East Asian' font slot (w:rFonts/@w:eastAsia) and
    will silently substitute a different font for CJK glyphs unless that
    slot is also set. Patching this at the style level covers every
    paragraph/run created via add_heading/add_paragraph without having to
    touch each individual run.
    """
    style_sizes = {
        "Normal": 9,
        "Heading 1": 9,
        "Heading 2": 9,
        "Heading 3": 9,
        "Heading 4": 9,
        "List Bullet": 9,
        "List Number": 9,
    }
    for name, size in style_sizes.items():
        try:
            style = doc.styles[name]
        except KeyError:
            continue
        style.font.name = font_name
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0x00, 0x00, 0x00)  ### djlee ### font color change
        rpr = style.element.get_or_add_rPr()
        rFonts = rpr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = OxmlElement('w:rFonts')
            rpr.append(rFonts)
        rFonts.set(qn('w:eastAsia'), font_name)
        rFonts.set(qn('w:ascii'), font_name)
        rFonts.set(qn('w:hAnsi'), font_name)


def _docx_add_inline(para, text: str):
    """Add text with bold/italic markdown markup to a docx paragraph."""
    # Split on **bold**, *italic*, ***bold-italic***
    parts = re.split(r'(\*\*\*.+?\*\*\*|\*\*.+?\*\*|\*.+?\*)', text)
    for part in parts:
        run = para.add_run()
        if re.match(r'^\*\*\*(.+)\*\*\*$', part):
            run.text = part[3:-3]
            run.bold = True
            run.italic = True
        elif re.match(r'^\*\*(.+)\*\*$', part):
            run.text = part[2:-2]
            run.bold = True
        elif re.match(r'^\*(.+)\*$', part):
            run.text = part[1:-1]
            run.italic = True
        else:
            run.text = re.sub(r'`([^`]+)`', r'\1', part)


def _docx_add_markdown(doc, md_text: str):
    """Render markdown blocks into the given python-docx Document."""
    for block in _parse_markdown(md_text or ""):
        btype = block["type"]

        if btype == "heading":
            lvl = min(block["level"] + 1, 4)  # shift down: md h1→docx h2
            p = doc.add_heading("", level=lvl)
            _docx_add_inline(p, block["text"])

        elif btype == "paragraph":
            p = doc.add_paragraph()
            _docx_add_inline(p, block["text"])

        elif btype == "ul":
            for item in block["items"]:
                p = doc.add_paragraph(style="List Bullet")
                _docx_add_inline(p, item)

        elif btype == "ol":
            for item in block["items"]:
                p = doc.add_paragraph(style="List Number")
                _docx_add_inline(p, item)

        elif btype == "table":
            headers = block["headers"]
            rows = block["rows"]
            ncols = len(headers)
            tbl = doc.add_table(rows=1 + len(rows), cols=ncols)
            tbl.style = "Table Grid"
            # Header row
            for ci, h in enumerate(headers):
                cell = tbl.rows[0].cells[ci]
                cell.text = _inline_to_plain(h)
                cell.paragraphs[0].runs[0].bold = True
            # Data rows
            for ri, row in enumerate(rows):
                for ci in range(ncols):
                    val = row[ci] if ci < len(row) else ""
                    tbl.rows[ri + 1].cells[ci].text = _inline_to_plain(val)
            doc.add_paragraph()

        elif btype == "code":
            p = doc.add_paragraph(block["text"])
            p.runs[0].font.name = "Courier New"

        elif btype == "hr":
            doc.add_paragraph("─" * 40)


def build_docx(meeting) -> bytes:
    doc = Document()
    _docx_set_style_font(doc)
    
    ### djlee ### 불필요한 내용 삭
    """
    # Title
    title_p = doc.add_heading(meeting.title, level=1)
    title_p.runs[0].font.color.rgb = RGBColor(0x1B, 0x14, 0x64)

    # Meta table
    meta = [
        ["일시", _fmt_date(meeting.meeting_date)],
        ["장소", meeting.location or "-"],
        ["참석자", meeting.attendees or "-"],
    ]
    tbl = doc.add_table(rows=len(meta), cols=2)
    tbl.style = "Table Grid"
    for i, (k, v) in enumerate(meta):
        tbl.rows[i].cells[0].text = k
        tbl.rows[i].cells[1].text = v
        tbl.rows[i].cells[0].paragraphs[0].runs[0].bold = True
    doc.add_paragraph()

    # Summary — render markdown
    doc.add_heading("주요 내용 요약", level=2)
    """
    _docx_add_markdown(doc, meeting.summary)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── PDF ──────────────────────────────────────────────────────────────────────

_ACCENT = colors.HexColor("#1B1464")

# Table styling — mirrors how Claude renders markdown tables in chat:
# a light gray header row with dark text, thin light borders, no stripes.
_TABLE_HEADER_BG = colors.HexColor("#F1F1F0")
_TABLE_HEADER_TEXT = colors.HexColor("#1F1E1D")
_TABLE_BORDER = colors.HexColor("#E0E0E0")
_TABLE_ROW_BG = colors.white

_PDF_PAGE_WIDTH = A4[0] - 4 * cm  # usable width (left+right margin = 4cm)


def _pdf_add_markdown(story: list, md_text: str, body_style, h2_style, font, font_bold):
    """Render parsed markdown blocks into a ReportLab story list."""
    styles = getSampleStyleSheet()

    bullet_style = ParagraphStyle(
        "MdBullet", parent=body_style,
        leftIndent=14, bulletIndent=0,
    )
    code_style = ParagraphStyle(
        "MdCode", parent=styles["Code"],
        fontName="Courier", fontSize=8, leading=12,
        backColor=colors.HexColor("#f4f4f4"),
        leftIndent=10, rightIndent=10,
        spaceBefore=4, spaceAfter=4,
    )

    tbl_header_style = TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), _TABLE_HEADER_BG),
        ("TEXTCOLOR", (0, 0), (-1, 0), _TABLE_HEADER_TEXT),
        ("FONTNAME", (0, 0), (-1, 0), font_bold),
        ("FONTNAME", (0, 1), (-1, -1), font),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, _TABLE_BORDER),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [_TABLE_ROW_BG]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ])

    heading_sizes = {1: 14, 2: 12, 3: 11, 4: 10, 5: 9, 6: 9}

    for block in _parse_markdown(md_text or ""):
        btype = block["type"]

        if btype == "heading":
            lvl = block["level"]
            hs = ParagraphStyle(
                f"MdH{lvl}", parent=h2_style,
                fontSize=heading_sizes.get(lvl, 10),
                spaceBefore=10, spaceAfter=4,
            )
            story.append(Paragraph(_inline_to_pdf(block["text"]), hs))

        elif btype == "paragraph":
            story.append(Paragraph(_inline_to_pdf(block["text"]), body_style))

        elif btype in ("ul", "ol"):
            for idx, item in enumerate(block["items"]):
                prefix = "•" if btype == "ul" else f"{idx + 1}."
                story.append(Paragraph(f"{prefix} {_inline_to_pdf(item)}", bullet_style))

        elif btype == "table":
            headers = block["headers"]
            rows = block["rows"]
            ncols = len(headers)

            def make_cell(text, is_header=False):
                fs = font_bold if is_header else font
                return Paragraph(
                    f'<font name="{fs}">{_inline_to_pdf(text)}</font>',
                    body_style,
                )

            data = [[make_cell(h, True) for h in headers]]
            for row in rows:
                data.append([make_cell(row[ci] if ci < len(row) else "") for ci in range(ncols)])

            col_w = _PDF_PAGE_WIDTH / ncols
            tbl = Table(data, colWidths=[col_w] * ncols, repeatRows=1)
            tbl.setStyle(tbl_header_style)
            story.append(tbl)
            story.append(Spacer(1, 0.3 * cm))

        elif btype == "code":
            story.append(Paragraph(block["text"].replace('\n', '<br/>'), code_style))

        elif btype == "hr":
            story.append(Spacer(1, 0.15 * cm))


def build_pdf(meeting) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2*cm, rightMargin=2*cm, topMargin=2*cm, bottomMargin=2*cm)
    styles = getSampleStyleSheet()

    h1 = ParagraphStyle("H1", parent=styles["Heading1"], textColor=_ACCENT, fontSize=16, spaceAfter=10, fontName=_KO_FONT)
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], textColor=_ACCENT, fontSize=12, spaceBefore=14, spaceAfter=6, fontName=_KO_FONT)
    body = ParagraphStyle("Body", parent=styles["Normal"], fontSize=10, leading=16, fontName=_KO_FONT)

    story = []
    story.append(Paragraph(_esc(meeting.title), h1))
    story.append(Spacer(1, 0.3*cm))

    meta_data = [
        ["일시", _fmt_date(meeting.meeting_date)],
        ["장소", meeting.location or "-"],
        ["참석자", meeting.attendees or "-"],
    ]
    meta_tbl = Table(meta_data, colWidths=[3*cm, 14*cm])
    meta_tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#ECEEF9")),
        ("FONTNAME", (0, 0), (-1, -1), _KO_FONT),
        ("FONTNAME", (0, 0), (0, -1), _KO_FONT_BOLD),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d2d5da")),
        ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, colors.HexColor("#fafafa")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
    ]))
    story.append(meta_tbl)
    story.append(Spacer(1, 0.5*cm))

    if meeting.summary:
        story.append(Paragraph("주요 내용 요약", h2))
        _pdf_add_markdown(story, meeting.summary, body, h2, _KO_FONT, _KO_FONT_BOLD)

    doc.build(story)
    return buf.getvalue()


# ── TXT ──────────────────────────────────────────────────────────────────────

def build_txt(meeting) -> bytes:
    lines = [
        f"{'='*60}",
        f"회의록: {meeting.title}",
        f"{'='*60}",
        f"일시    : {_fmt_date(meeting.meeting_date)}",
        f"장소    : {meeting.location or '-'}",
        f"참석자  : {meeting.attendees or '-'}",
        "",
        "[주요 내용 요약]",
        meeting.summary or "",
        "",
    ]
    return "\n".join(lines).encode("utf-8")
